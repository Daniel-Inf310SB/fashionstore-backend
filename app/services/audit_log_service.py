from __future__ import annotations

import json
import math

from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from sqlalchemy import (
    func,
    or_,
    select,
)

from sqlalchemy.orm import (
    Session,
    selectinload,
)

from app.models.audit_log import (
    AuditLog,
)

from app.models.user import (
    User,
)


class AuditLogService:

    # =====================================================
    # REGISTRAR EVENTO
    # =====================================================

    @staticmethod
    def log(
        db: Session,
        *,
        user_id: int | None,
        action: str,
        module: str,
        entity_type: str | None = None,
        entity_id: int | None = None,
        description: str | None = None,
        old_values:
            dict[str, Any] | None = None,
        new_values:
            dict[str, Any] | None = None,
        ip_address: str | None = None,
        user_agent: str | None = None,
        status: str = "SUCCESS",
    ) -> AuditLog:

        audit_log = AuditLog(
            user_id=
                user_id,

            action=
                action.strip().upper(),

            module=
                module.strip().upper(),

            entity_type=
                (
                    entity_type.strip()
                    if entity_type
                    else None
                ),

            entity_id=
                entity_id,

            description=
                description,

            old_values=
                old_values,

            new_values=
                new_values,

            ip_address=
                ip_address,

            user_agent=
                user_agent,

            status=
                status.strip().upper(),
        )

        db.add(
            audit_log
        )

        db.flush()

        return audit_log


    # =====================================================
    # CONSTRUIR FILTROS
    # =====================================================

    @staticmethod
    def _build_filters(
        *,
        search: str | None = None,
        user_id: int | None = None,
        action: str | None = None,
        module: str | None = None,
        entity_type: str | None = None,
        entity_id: int | None = None,
        status: str | None = None,
        date_from:
            datetime | None = None,
        date_to:
            datetime | None = None,
    ) -> list:

        filters = []


        # =================================================
        # BÚSQUEDA GENERAL
        # =================================================

        if (
            search
            and search.strip()
        ):

            value = (
                f"%{search.strip()}%"
            )

            filters.append(
                or_(
                    AuditLog.action.ilike(
                        value
                    ),

                    AuditLog.module.ilike(
                        value
                    ),

                    AuditLog.entity_type.ilike(
                        value
                    ),

                    AuditLog.description.ilike(
                        value
                    ),

                    AuditLog.ip_address.ilike(
                        value
                    ),

                    AuditLog.status.ilike(
                        value
                    ),

                    AuditLog.user.has(
                        or_(
                            User.username.ilike(
                                value
                            ),

                            User.first_name.ilike(
                                value
                            ),

                            User.last_name.ilike(
                                value
                            ),

                            User.email.ilike(
                                value
                            ),
                        )
                    ),
                )
            )


        # =================================================
        # USUARIO
        # =================================================

        if user_id is not None:

            filters.append(
                AuditLog.user_id
                == user_id
            )


        # =================================================
        # ACCIÓN
        # =================================================

        if (
            action
            and action.strip()
        ):

            filters.append(
                AuditLog.action
                == action
                .strip()
                .upper()
            )


        # =================================================
        # MÓDULO
        # =================================================

        if (
            module
            and module.strip()
        ):

            filters.append(
                AuditLog.module
                == module
                .strip()
                .upper()
            )


        # =================================================
        # ENTIDAD
        # =================================================

        if (
            entity_type
            and entity_type.strip()
        ):

            filters.append(
                func.lower(
                    AuditLog.entity_type
                )
                == entity_type
                .strip()
                .lower()
            )


        # =================================================
        # ID ENTIDAD
        # =================================================

        if entity_id is not None:

            filters.append(
                AuditLog.entity_id
                == entity_id
            )


        # =================================================
        # ESTADO
        # =================================================

        if (
            status
            and status.strip()
        ):

            filters.append(
                AuditLog.status
                == status
                .strip()
                .upper()
            )


        # =================================================
        # FECHA DESDE
        # =================================================

        if date_from is not None:

            filters.append(
                AuditLog.created_at
                >= date_from
            )


        # =================================================
        # FECHA HASTA
        # =================================================

        if date_to is not None:

            filters.append(
                AuditLog.created_at
                <= date_to
            )


        return filters


    # =====================================================
    # OBTENER COLUMNA DE ORDENAMIENTO
    # =====================================================

    @staticmethod
    def _get_ordering(
        sort_by: str,
        sort_order: str,
    ):

        sort_columns = {
            "id":
                AuditLog.id,

            "created_at":
                AuditLog.created_at,

            "action":
                AuditLog.action,

            "module":
                AuditLog.module,

            "status":
                AuditLog.status,

            "user_id":
                AuditLog.user_id,

            "entity_type":
                AuditLog.entity_type,

            "entity_id":
                AuditLog.entity_id,
        }


        sort_column = (
            sort_columns.get(
                sort_by,
                AuditLog.created_at,
            )
        )


        if (
            sort_order.lower()
            == "asc"
        ):

            return sort_column.asc()


        return sort_column.desc()


    # =====================================================
    # LISTAR BITÁCORA
    # =====================================================

    @staticmethod
    def list_audit_logs(
        db: Session,

        page: int = 1,

        page_size: int = 10,

        search: str | None = None,

        user_id: int | None = None,

        action: str | None = None,

        module: str | None = None,

        entity_type: str | None = None,

        entity_id: int | None = None,

        status: str | None = None,

        date_from:
            datetime | None = None,

        date_to:
            datetime | None = None,

        sort_by: str = "created_at",

        sort_order: str = "desc",
    ) -> dict:

        page = max(
            page,
            1,
        )

        page_size = min(
            max(
                page_size,
                1,
            ),
            100,
        )


        filters = (
            AuditLogService
            ._build_filters(
                search=
                    search,

                user_id=
                    user_id,

                action=
                    action,

                module=
                    module,

                entity_type=
                    entity_type,

                entity_id=
                    entity_id,

                status=
                    status,

                date_from=
                    date_from,

                date_to=
                    date_to,
            )
        )


        total = (
            db.scalar(
                select(
                    func.count(
                        AuditLog.id
                    )
                )
                .where(
                    *filters
                )
            )
            or 0
        )


        ordering = (
            AuditLogService
            ._get_ordering(
                sort_by=
                    sort_by,

                sort_order=
                    sort_order,
            )
        )


        statement = (
            select(
                AuditLog
            )
            .options(
                selectinload(
                    AuditLog.user
                )
            )
            .where(
                *filters
            )
            .order_by(
                ordering
            )
            .offset(
                (page - 1)
                * page_size
            )
            .limit(
                page_size
            )
        )


        audit_logs = (
            db.scalars(
                statement
            )
            .all()
        )


        total_pages = (
            math.ceil(
                total
                / page_size
            )
            if total > 0
            else 0
        )


        return {
            "items":
                audit_logs,

            "page":
                page,

            "page_size":
                page_size,

            "total":
                total,

            "total_pages":
                total_pages,
        }


    # =====================================================
    # OBTENER REGISTROS PARA EXPORTACIÓN
    # =====================================================

    @staticmethod
    def _get_logs_for_export(
        db: Session,
        *,
        search: str | None = None,
        user_id: int | None = None,
        action: str | None = None,
        module: str | None = None,
        entity_type: str | None = None,
        entity_id: int | None = None,
        status: str | None = None,
        date_from:
            datetime | None = None,
        date_to:
            datetime | None = None,
        sort_by: str = "created_at",
        sort_order: str = "desc",
    ) -> list[AuditLog]:

        filters = (
            AuditLogService
            ._build_filters(
                search=
                    search,

                user_id=
                    user_id,

                action=
                    action,

                module=
                    module,

                entity_type=
                    entity_type,

                entity_id=
                    entity_id,

                status=
                    status,

                date_from=
                    date_from,

                date_to=
                    date_to,
            )
        )


        ordering = (
            AuditLogService
            ._get_ordering(
                sort_by=
                    sort_by,

                sort_order=
                    sort_order,
            )
        )


        statement = (
            select(
                AuditLog
            )
            .options(
                selectinload(
                    AuditLog.user
                )
            )
            .where(
                *filters
            )
            .order_by(
                ordering
            )
        )


        return list(
            db.scalars(
                statement
            )
            .all()
        )


    # =====================================================
    # NOMBRE DEL USUARIO
    # =====================================================

    @staticmethod
    def _user_name(
        audit_log: AuditLog,
    ) -> str:

        if audit_log.user is None:

            return "Sistema"


        first_name = (
            audit_log.user.first_name
            or ""
        ).strip()

        last_name = (
            audit_log.user.last_name
            or ""
        ).strip()


        full_name = (
            f"{first_name} {last_name}"
            .strip()
        )


        return (
            full_name
            or audit_log.user.username
            or audit_log.user.email
            or "Usuario"
        )


    # =====================================================
    # FORMATEAR FECHA
    # =====================================================

    @staticmethod
    def _format_datetime(
        value: datetime | None,
    ) -> str:

        if value is None:

            return ""


        return value.strftime(
            "%d/%m/%Y %H:%M:%S"
        )


    # =====================================================
    # FORMATEAR JSON
    # =====================================================

    @staticmethod
    def _format_json(
        value: dict[str, Any] | None,
    ) -> str:

        if not value:

            return ""


        return json.dumps(
            value,
            ensure_ascii=False,
            default=str,
        )


    # =====================================================
    # EXPORTAR EXCEL
    # =====================================================

    @staticmethod
    def export_excel(
        db: Session,
        *,
        search: str | None = None,
        user_id: int | None = None,
        action: str | None = None,
        module: str | None = None,
        entity_type: str | None = None,
        entity_id: int | None = None,
        status: str | None = None,
        date_from:
            datetime | None = None,
        date_to:
            datetime | None = None,
        sort_by: str = "created_at",
        sort_order: str = "desc",
    ) -> BytesIO:

        audit_logs = (
            AuditLogService
            ._get_logs_for_export(
                db=db,

                search=
                    search,

                user_id=
                    user_id,

                action=
                    action,

                module=
                    module,

                entity_type=
                    entity_type,

                entity_id=
                    entity_id,

                status=
                    status,

                date_from=
                    date_from,

                date_to=
                    date_to,

                sort_by=
                    sort_by,

                sort_order=
                    sort_order,
            )
        )


        workbook = Workbook()

        worksheet = workbook.active

        worksheet.title = (
            "Bitácora de auditoría"
        )


        # =================================================
        # TÍTULO
        # =================================================

        worksheet.merge_cells(
            "A1:L1"
        )

        title_cell = (
            worksheet["A1"]
        )

        title_cell.value = (
            "FashionStore - Bitácora de auditoría"
        )

        title_cell.font = Font(
            bold=True,
            size=16,
            color="FFFFFF",
        )

        title_cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1E293B",
        )

        title_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        worksheet.row_dimensions[
            1
        ].height = 28


        # =================================================
        # FECHA DE GENERACIÓN
        # =================================================

        worksheet.merge_cells(
            "A2:L2"
        )

        worksheet["A2"] = (
            "Generado: "
            f"{datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
        )

        worksheet["A2"].alignment = Alignment(
            horizontal="center",
        )


        # =================================================
        # CABECERAS
        # =================================================

        headers = [
            "ID",
            "Fecha",
            "Usuario",
            "Correo",
            "Acción",
            "Módulo",
            "Entidad",
            "ID entidad",
            "Descripción",
            "Estado",
            "IP",
            "Cambios",
        ]


        for column, header in enumerate(
            headers,
            start=1,
        ):

            cell = worksheet.cell(
                row=4,
                column=column,
                value=header,
            )

            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="4F46E5",
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )


        # =================================================
        # DATOS
        # =================================================

        for row_number, audit_log in enumerate(
            audit_logs,
            start=5,
        ):

            user_email = (
                audit_log.user.email
                if audit_log.user
                else ""
            )


            old_values = (
                AuditLogService
                ._format_json(
                    audit_log.old_values
                )
            )

            new_values = (
                AuditLogService
                ._format_json(
                    audit_log.new_values
                )
            )


            changes = ""

            if old_values:

                changes += (
                    f"Antes: {old_values}"
                )


            if new_values:

                if changes:

                    changes += "\n"

                changes += (
                    f"Después: {new_values}"
                )


            values = [
                audit_log.id,

                AuditLogService
                ._format_datetime(
                    audit_log.created_at
                ),

                AuditLogService
                ._user_name(
                    audit_log
                ),

                user_email,

                audit_log.action,

                audit_log.module,

                audit_log.entity_type
                or "",

                audit_log.entity_id
                or "",

                audit_log.description
                or "",

                audit_log.status,

                audit_log.ip_address
                or "",

                changes,
            ]


            for column, value in enumerate(
                values,
                start=1,
            ):

                cell = worksheet.cell(
                    row=
                        row_number,

                    column=
                        column,

                    value=
                        value,
                )

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )


        # =================================================
        # ANCHOS
        # =================================================

        widths = {
            "A": 8,
            "B": 21,
            "C": 24,
            "D": 30,
            "E": 28,
            "F": 18,
            "G": 18,
            "H": 12,
            "I": 40,
            "J": 14,
            "K": 18,
            "L": 60,
        }


        for column, width in (
            widths.items()
        ):

            worksheet.column_dimensions[
                column
            ].width = width


        worksheet.freeze_panes = (
            "A5"
        )


        # =================================================
        # AUTOFILTRO
        # =================================================

        if audit_logs:

            worksheet.auto_filter.ref = (
                f"A4:L{4 + len(audit_logs)}"
            )


        output = BytesIO()

        workbook.save(
            output
        )

        output.seek(
            0
        )

        return output


    # =====================================================
    # EXPORTAR PDF
    # =====================================================

    @staticmethod
    def export_pdf(
        db: Session,
        *,
        search: str | None = None,
        user_id: int | None = None,
        action: str | None = None,
        module: str | None = None,
        entity_type: str | None = None,
        entity_id: int | None = None,
        status: str | None = None,
        date_from:
            datetime | None = None,
        date_to:
            datetime | None = None,
        sort_by: str = "created_at",
        sort_order: str = "desc",
    ) -> BytesIO:

        audit_logs = (
            AuditLogService
            ._get_logs_for_export(
                db=db,

                search=
                    search,

                user_id=
                    user_id,

                action=
                    action,

                module=
                    module,

                entity_type=
                    entity_type,

                entity_id=
                    entity_id,

                status=
                    status,

                date_from=
                    date_from,

                date_to=
                    date_to,

                sort_by=
                    sort_by,

                sort_order=
                    sort_order,
            )
        )


        output = BytesIO()


        document = SimpleDocTemplate(
            output,

            pagesize=
                landscape(
                    A4
                ),

            rightMargin=
                12 * mm,

            leftMargin=
                12 * mm,

            topMargin=
                12 * mm,

            bottomMargin=
                12 * mm,
        )


        styles = (
            getSampleStyleSheet()
        )


        title_style = ParagraphStyle(
            "AuditTitle",

            parent=
                styles["Title"],

            alignment=
                TA_CENTER,

            fontSize=
                17,

            leading=
                21,

            spaceAfter=
                8,
        )


        subtitle_style = ParagraphStyle(
            "AuditSubtitle",

            parent=
                styles["Normal"],

            alignment=
                TA_CENTER,

            fontSize=
                9,

            textColor=
                colors.HexColor(
                    "#64748B"
                ),

            spaceAfter=
                14,
        )


        cell_style = ParagraphStyle(
            "AuditCell",

            parent=
                styles["Normal"],

            fontSize=
                7,

            leading=
                9,
        )


        header_style = ParagraphStyle(
            "AuditHeader",

            parent=
                styles["Normal"],

            fontSize=
                7,

            leading=
                9,

            textColor=
                colors.white,

            alignment=
                TA_CENTER,
        )


        story = []


        story.append(
            Paragraph(
                "FashionStore - Bitácora de auditoría",
                title_style,
            )
        )


        story.append(
            Paragraph(
                (
                    "Generado: "
                    f"{datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
                    " | "
                    f"Registros: {len(audit_logs)}"
                ),
                subtitle_style,
            )
        )


        story.append(
            Spacer(
                1,
                4,
            )
        )


        headers = [
            "Fecha",
            "Usuario",
            "Acción",
            "Módulo",
            "Entidad",
            "Descripción",
            "Estado",
            "IP",
        ]


        table_data = [
            [
                Paragraph(
                    header,
                    header_style,
                )
                for header in headers
            ]
        ]


        for audit_log in audit_logs:

            entity = (
                audit_log.entity_type
                or "—"
            )


            if (
                audit_log.entity_id
                is not None
            ):

                entity += (
                    f" #{audit_log.entity_id}"
                )


            row = [
                AuditLogService
                ._format_datetime(
                    audit_log.created_at
                ),

                AuditLogService
                ._user_name(
                    audit_log
                ),

                audit_log.action,

                audit_log.module,

                entity,

                audit_log.description
                or "—",

                audit_log.status,

                audit_log.ip_address
                or "—",
            ]


            table_data.append(
                [
                    Paragraph(
                        str(value),
                        cell_style,
                    )
                    for value in row
                ]
            )


        table = Table(
            table_data,

            repeatRows=
                1,

            colWidths=[
                28 * mm,
                36 * mm,
                35 * mm,
                25 * mm,
                30 * mm,
                60 * mm,
                20 * mm,
                27 * mm,
            ],
        )


        table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor(
                            "#4F46E5"
                        ),
                    ),

                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white,
                    ),

                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),

                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.25,
                        colors.HexColor(
                            "#CBD5E1"
                        ),
                    ),

                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [
                            colors.white,
                            colors.HexColor(
                                "#F8FAFC"
                            ),
                        ],
                    ),

                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        4,
                    ),

                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        4,
                    ),

                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),

                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                ]
            )
        )


        story.append(
            table
        )


        document.build(
            story
        )


        output.seek(
            0
        )

        return output


    # =====================================================
    # OBTENER DETALLE
    # =====================================================

    @staticmethod
    def get_audit_log(
        db: Session,

        audit_log_id: int,
    ) -> AuditLog:

        statement = (
            select(
                AuditLog
            )
            .options(
                selectinload(
                    AuditLog.user
                )
            )
            .where(
                AuditLog.id
                == audit_log_id
            )
        )


        audit_log = (
            db.scalar(
                statement
            )
        )


        if audit_log is None:

            raise LookupError(
                "Registro de bitácora no encontrado"
            )


        return audit_log