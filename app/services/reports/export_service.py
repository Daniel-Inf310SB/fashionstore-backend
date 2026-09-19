from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape


class ReportExportService:
    MEDIA_TYPES = {
        "pdf": "application/pdf",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }

    @staticmethod
    def _value(value: Any) -> str:
        if value is None:
            return "—"
        if isinstance(value, float):
            return f"{value:,.2f}"
        if isinstance(value, bool):
            return "Sí" if value else "No"
        return str(value)

    @staticmethod
    def _filter_lines(report: dict[str, Any]) -> list[tuple[str, str]]:
        filters = report.get("filters") or {}
        if not filters:
            return [("Filtros", "Sin filtros")]
        return [(key.replace("_", " ").title(), ReportExportService._value(value)) for key, value in filters.items()]

    @staticmethod
    def _summary_lines(report: dict[str, Any]) -> list[tuple[str, str]]:
        summary = report.get("summary") or {}
        rows: list[tuple[str, str]] = []
        for key, value in summary.items():
            label = key.replace("_", " ").title()
            if isinstance(value, dict):
                text = ", ".join(f"{k}: {v}" for k, v in value.items()) or "—"
            else:
                text = ReportExportService._value(value)
            rows.append((label, text))
        return rows

    @staticmethod
    def export(report: dict[str, Any], export_format: str) -> tuple[bytes, str, str]:
        normalized = export_format.lower().strip()
        if normalized == "pdf":
            content = ReportExportService.to_pdf(report)
        elif normalized == "xlsx":
            content = ReportExportService.to_excel(report)
        elif normalized == "docx":
            content = ReportExportService.to_word(report)
        else:
            raise ValueError("Formato no soportado. Use pdf, xlsx o docx.")

        return content, ReportExportService.MEDIA_TYPES[normalized], normalized

    @staticmethod
    def to_pdf(report: dict[str, Any]) -> bytes:
        try:
            from reportlab.lib import colors
            from reportlab.lib.enums import TA_CENTER
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError as exc:
            raise RuntimeError("Falta instalar reportlab: pip install reportlab") from exc

        output = BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=10 * mm,
            leftMargin=10 * mm,
            topMargin=10 * mm,
            bottomMargin=10 * mm,
            title=report["title"],
            author="FashionStore",
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "ReportTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            fontSize=18,
            leading=22,
            spaceAfter=8,
        )
        small = ParagraphStyle("ReportSmall", parent=styles["BodyText"], fontSize=7.5, leading=9)
        meta = ParagraphStyle("ReportMeta", parent=styles["BodyText"], fontSize=8.5, leading=10)

        story: list[Any] = [
            Paragraph("FashionStore", title_style),
            Paragraph(escape(report["title"]), styles["Heading2"]),
            Paragraph(f"Generado: {escape(str(report['generated_at']))}", meta),
            Spacer(1, 5 * mm),
        ]

        filter_data = [["Filtros aplicados", "Valor"]] + [list(row) for row in ReportExportService._filter_lines(report)]
        filter_table = Table(filter_data, colWidths=[55 * mm, 200 * mm], repeatRows=1)
        filter_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#111111")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D9D9D9")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F7F7")]),
        ]))
        story += [filter_table, Spacer(1, 5 * mm)]

        summary_data = [["Resumen", "Valor"]] + [list(row) for row in ReportExportService._summary_lines(report)]
        summary_table = Table(summary_data, colWidths=[75 * mm, 180 * mm], repeatRows=1)
        summary_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E91E63")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D9D9D9")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story += [summary_table, Spacer(1, 6 * mm)]

        columns = report.get("columns") or []
        header = [Paragraph(f"<b>{escape(col['label'])}</b>", small) for col in columns]
        data: list[list[Any]] = [header]
        for row in report.get("rows") or []:
            data.append([
                Paragraph(escape(ReportExportService._value(row.get(col["key"]))), small)
                for col in columns
            ])

        if len(data) == 1:
            data.append([Paragraph("Sin resultados", small)] + [""] * max(0, len(columns) - 1))

        page_width = landscape(A4)[0] - 20 * mm
        col_width = page_width / max(1, len(columns))
        table = Table(data, colWidths=[col_width] * len(columns), repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#111111")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D0D0")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8F8F8")]),
        ]))
        story.append(table)
        doc.build(story)
        return output.getvalue()

    @staticmethod
    def to_excel(report: dict[str, Any]) -> bytes:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise RuntimeError("Falta instalar openpyxl: pip install openpyxl") from exc

        wb = Workbook()
        ws = wb.active
        ws.title = "Reporte"

        ws["A1"] = "FashionStore"
        ws["A1"].font = Font(size=18, bold=True)
        ws["A2"] = report["title"]
        ws["A2"].font = Font(size=14, bold=True)
        ws["A3"] = f"Generado: {report['generated_at']}"

        row_index = 5
        ws.cell(row=row_index, column=1, value="FILTROS APLICADOS").font = Font(bold=True)
        row_index += 1
        for label, value in ReportExportService._filter_lines(report):
            ws.cell(row=row_index, column=1, value=label)
            ws.cell(row=row_index, column=2, value=value)
            row_index += 1

        row_index += 1
        ws.cell(row=row_index, column=1, value="RESUMEN").font = Font(bold=True)
        row_index += 1
        for label, value in ReportExportService._summary_lines(report):
            ws.cell(row=row_index, column=1, value=label)
            ws.cell(row=row_index, column=2, value=value)
            row_index += 1

        row_index += 2
        table_header_row = row_index
        columns = report.get("columns") or []
        header_fill = PatternFill("solid", fgColor="111111")
        for col_index, column in enumerate(columns, start=1):
            cell = ws.cell(row=row_index, column=col_index, value=column["label"])
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for data_row in report.get("rows") or []:
            row_index += 1
            for col_index, column in enumerate(columns, start=1):
                value = data_row.get(column["key"])
                cell = ws.cell(row=row_index, column=col_index, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        ws.freeze_panes = f"A{table_header_row + 1}"
        if columns:
            ws.auto_filter.ref = f"A{table_header_row}:{get_column_letter(len(columns))}{row_index}"

        for col_index, column in enumerate(columns, start=1):
            max_len = len(column["label"])
            for r in range(table_header_row + 1, row_index + 1):
                value = ws.cell(r, col_index).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            ws.column_dimensions[get_column_letter(col_index)].width = min(max(max_len + 2, 12), 38)

        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def to_word(report: dict[str, Any]) -> bytes:
        try:
            from docx import Document
            from docx.enum.section import WD_ORIENT
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Inches, Pt
        except ImportError as exc:
            raise RuntimeError("Falta instalar python-docx: pip install python-docx") from exc

        doc = Document()
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.45)
        section.left_margin = Inches(0.45)
        section.right_margin = Inches(0.45)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("FashionStore")
        run.bold = True
        run.font.size = Pt(18)

        heading = doc.add_heading(report["title"], level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Generado: {report['generated_at']}")

        doc.add_heading("Filtros aplicados", level=2)
        filters_table = doc.add_table(rows=1, cols=2)
        filters_table.style = "Table Grid"
        filters_table.rows[0].cells[0].text = "Filtro"
        filters_table.rows[0].cells[1].text = "Valor"
        for label, value in ReportExportService._filter_lines(report):
            cells = filters_table.add_row().cells
            cells[0].text = label
            cells[1].text = value

        doc.add_heading("Resumen", level=2)
        summary_table = doc.add_table(rows=1, cols=2)
        summary_table.style = "Table Grid"
        summary_table.rows[0].cells[0].text = "Indicador"
        summary_table.rows[0].cells[1].text = "Valor"
        for label, value in ReportExportService._summary_lines(report):
            cells = summary_table.add_row().cells
            cells[0].text = label
            cells[1].text = value

        doc.add_heading("Detalle", level=2)
        columns = report.get("columns") or []
        detail = doc.add_table(rows=1, cols=max(1, len(columns)))
        detail.style = "Table Grid"
        if columns:
            for index, column in enumerate(columns):
                detail.rows[0].cells[index].text = column["label"]
            for data_row in report.get("rows") or []:
                cells = detail.add_row().cells
                for index, column in enumerate(columns):
                    cells[index].text = ReportExportService._value(data_row.get(column["key"]))
        else:
            detail.rows[0].cells[0].text = "Sin datos"

        for table in [filters_table, summary_table, detail]:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(7.5)

        output = BytesIO()
        doc.save(output)
        return output.getvalue()
